from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo


SCHEMA_VERSION = "0.1"
TIMEZONE = ZoneInfo("Europe/Lisbon")

PROJECT_ROOT = Path(r"G:\My Drive\Claude_PRJ")
AGENDA_ROOT = PROJECT_ROOT / "Agenda"
ROUTINE_DIR = AGENDA_ROOT / "X_Rotinas_Python" / "Painel_Hff"
TEMPLATES_DIR = AGENDA_ROOT / "Templates"
HFF_TEMPLATE_PATH = TEMPLATES_DIR / "hff" / "HFF_TEMPLATE_v6_render.html"
HFF_RENDER_PATH = TEMPLATES_DIR / "hff" / "render_v4.py"
BO_TEMPLATE_PATH = TEMPLATES_DIR / "bo" / "BO_TEMPLATE_v4_render.html"
BO_RENDER_PATH = TEMPLATES_DIR / "bo" / "render_v3.py"
REPO_DIR = PROJECT_ROOT / "personal-system"
REPO_HFF_PATH = REPO_DIR / "agendas" / "Hff" / "index.html"
REPO_BO_PATH = REPO_DIR / "agendas" / "BO" / "index.html"

INSTRUCOES_PATH = AGENDA_ROOT / "INSTRUCOES_PAINEL_HFF.md"
PUBLICACAO_PATH = AGENDA_ROOT / "ROTINA_PUBLICACAO.md"
SYSTEM_PROMPT_PATH = AGENDA_ROOT / "SYSTEM_PROMPT_FINAL.md"

ESPELHO_HFF_ID = "1culq10MksUxSd05P1_ejqOEzaoQrFLKT-vP0brqp7iU"
WINDOW_DAYS = 28

WEEKDAYS_PT = {
    0: "segunda",
    1: "terca",
    2: "quarta",
    3: "quinta",
    4: "sexta",
    5: "sabado",
    6: "domingo",
}
MONTHS_PT = {
    "janeiro": 1,
    "jan": 1,
    "fevereiro": 2,
    "fev": 2,
    "marco": 3,
    "mar": 3,
    "abril": 4,
    "abr": 4,
    "maio": 5,
    "mai": 5,
    "junho": 6,
    "jun": 6,
    "julho": 7,
    "jul": 7,
    "agosto": 8,
    "ago": 8,
    "setembro": 9,
    "set": 9,
    "outubro": 10,
    "out": 10,
    "novembro": 11,
    "nov": 11,
    "dezembro": 12,
    "dez": 12,
}
TEAM = {
    "IF": "Isabel Franca",
    "RC": "Rafael Correia",
    "RR": "Rodrigo Roquette",
    "A": "Afonso",
}
NAME_TO_INITIALS = {
    "isabel": "IF",
    "isabel franca": "IF",
    "rafael": "RC",
    "rafael correia": "RC",
    "rodrigo": "RR",
    "rodrigo roquette": "RR",
    "afonso": "A",
}
ABSENCE_RE = re.compile(r"^(?P<motive>[FCM])(?P<who>if|rc|rr|a)$", re.IGNORECASE)
DAY_OFF_RE = re.compile(r"^\*(?P<who>if|rc|rr|a)$", re.IGNORECASE)


class SimpleCell:
    def __init__(self, row: int, column: int) -> None:
        self.coordinate = f"{column_letter(column)}{row}"


class ValueSheet:
    def __init__(self, title: str, rows: list[list[Any]]) -> None:
        self.title = title
        self._rows = rows

    def iter_rows(self, values_only: bool = True) -> list[list[Any]]:
        return self._rows

    def cell(self, row: int, column: int) -> SimpleCell:
        return SimpleCell(row, column)


class ValueWorkbook:
    def __init__(self, sheets: dict[str, list[list[Any]]]) -> None:
        self._sheets = {name: ValueSheet(name, rows) for name, rows in sheets.items()}
        self.sheetnames = list(self._sheets.keys())

    def __getitem__(self, name: str) -> ValueSheet:
        return self._sheets[name]


def column_letter(column: int) -> str:
    result = ""
    while column:
        column, rem = divmod(column - 1, 26)
        result = chr(65 + rem) + result
    return result


def now_lisbon() -> datetime:
    return datetime.now(TIMEZONE)


def file_info(path: Path | None) -> dict[str, Any]:
    if path is None:
        return {"exists": False, "path": None}
    exists = path.exists()
    info: dict[str, Any] = {"exists": exists, "path": str(path)}
    if exists:
        stat = path.stat()
        info.update(
            {
                "name": path.name,
                "size_bytes": stat.st_size,
                "mtime": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
            }
        )
    return info


def norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text.strip().lower())


def compact_norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9*]+", "", norm(value))


def parse_date(value: Any, default_year: int | None = None) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        if 20000 <= float(value) <= 60000:
            return date(1899, 12, 30) + timedelta(days=int(value))
        return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    # Formato visto na coluna "Data Cirurgia" quando a celula chega como texto
    # em vez de data nativa: "qua., 23/09/26" (dia da semana abreviado + data).
    # Um doente real ficou silenciosamente de fora da lista por causa disto.
    m = re.match(r"^[a-zA-Zà-úÀ-Ú]{3}\.?,?\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", text)
    if m:
        day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if year < 100:
            year += 2000
        try:
            return date(year, month, day)
        except ValueError:
            return None
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})$", text)
    if m and default_year:
        try:
            return date(default_year, int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = norm(value)
    return text in {"true", "t", "1", "sim", "yes", "y", "verdadeiro"}


def parse_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def age_fmt(value: Any) -> str:
    age = parse_float(value)
    if age is None:
        return "s/d"
    if age <= 0:
        return "< 1 mes"
    if age < 1:
        months = max(1, round(age * 12))
        return f"{months} mes" if months == 1 else f"{months} meses"
    return f"{age:.1f}a".replace(".", ",")


def load_workbook(path: Path) -> Any:
    try:
        from openpyxl import load_workbook as openpyxl_load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl nao esta instalado no ambiente Python") from exc
    return openpyxl_load_workbook(path, data_only=True, read_only=True)


def row_values(ws: Any) -> list[list[Any]]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def sheet_by_name(wb: Any, candidates: list[str]) -> Any | None:
    normalized = {norm(name): name for name in wb.sheetnames}
    for candidate in candidates:
        real = normalized.get(norm(candidate))
        if real:
            return wb[real]
    return None


def find_header(rows: list[list[Any]], aliases: dict[str, list[str]], minimum: int) -> tuple[int | None, dict[str, int]]:
    alias_lookup: dict[str, str] = {}
    for field, values in aliases.items():
        for value in values:
            alias_lookup[compact_norm(value)] = field
    best_idx: int | None = None
    best_map: dict[str, int] = {}
    for idx, row in enumerate(rows[:20]):
        current: dict[str, int] = {}
        for col_idx, cell in enumerate(row):
            field = alias_lookup.get(compact_norm(cell))
            if field and field not in current:
                current[field] = col_idx
        if len(current) > len(best_map):
            best_idx = idx
            best_map = current
        if len(current) >= minimum:
            return idx, current
    return best_idx, best_map


def cell(row: list[Any], mapping: dict[str, int], field: str) -> Any:
    idx = mapping.get(field)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def initials_from_name(value: Any) -> str | None:
    text = norm(value)
    if not text:
        return None
    if text.upper() in TEAM:
        return text.upper()
    for key, initials in NAME_TO_INITIALS.items():
        if key in text:
            return initials
    return None


def motive_label(code: str) -> str:
    return {"F": "Ferias", "C": "Curso", "M": "Madeira"}.get(code.upper(), code)


def inspect_templates_and_repo() -> dict[str, Any]:
    return {
        "instrucoes_hff": file_info(INSTRUCOES_PATH),
        "rotina_publicacao": file_info(PUBLICACAO_PATH),
        "system_prompt": file_info(SYSTEM_PROMPT_PATH),
        "hff_template": file_info(HFF_TEMPLATE_PATH),
        "hff_render": file_info(HFF_RENDER_PATH),
        "bo_template": file_info(BO_TEMPLATE_PATH),
        "bo_render": file_info(BO_RENDER_PATH),
        "repo": file_info(REPO_DIR),
        "repo_hff_index": file_info(REPO_HFF_PATH),
        "repo_bo_index": file_info(REPO_BO_PATH),
    }


def parse_cirurgias(ws: Any, today: date, warnings: list[str], errors: list[str]) -> dict[str, Any]:
    aliases = {
        "data": ["Data Cirurgia"],
        "periodo": ["Periodo", "Periodo (M = manha / S = SIGIC-tarde)", "Periodo (Manha/Sigic (tarde))"],
        "processo": ["Processo"],
        "nome": ["Nome doente", "Nome"],
        "fdr": ["FDR", "FDR (Fora da Rotina)"],
        "tempo_espera": ["Tempo em Lista de Espera"],
        "idade": ["Idade"],
        "procedimento": ["Procedimento"],
        "status": ["Status Agendamento"],
        "obs": ["Observacao", "Observação"],
        "if_ausente": ["Previsao IF ausente", "Previsão IF ausente", "Previsao de IF estar ausente no dia"],
        "rc_ausente": ["Previsao RC ausente", "Previsão RC ausente", "Previsao de RC estar ausente no dia"],
        "rr_ausente": ["Previsao RR ausente", "Previsão RR ausente", "Previsao de RR estar ausente no dia"],
    }
    rows = row_values(ws)
    header_idx, mapping = find_header(rows, aliases, minimum=8)
    missing_headers = sorted(set(aliases) - set(mapping))
    if header_idx is None or len(mapping) < 8:
        errors.append("Cirurgias: cabecalho obrigatorio nao encontrado")
        return {"ok": False, "headers_found": sorted(mapping), "missing_headers": missing_headers}
    if missing_headers:
        warnings.append("Cirurgias: cabecalhos ausentes/parciais: " + ", ".join(missing_headers))

    end = today + timedelta(days=WINDOW_DAYS - 1)
    raw_rows: list[dict[str, Any]] = []
    sessions: dict[tuple[str, str], dict[str, Any]] = {}
    anomalies: list[dict[str, Any]] = []
    process_keys: list[tuple[str, str, str]] = []

    for excel_row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        raw_date_cell = cell(row, mapping, "data")
        surgery_date = parse_date(raw_date_cell, default_year=today.year)
        if surgery_date is None:
            if str(raw_date_cell or "").strip():
                # Nunca cair fora da lista em silencio: uma data ilegivel e um
                # doente a menos que ninguem repara ter desaparecido.
                anomalies.append({
                    "type": "data_ilegivel",
                    "row": excel_row_idx,
                    "raw": str(raw_date_cell),
                    "processo": str(cell(row, mapping, "processo") or "").strip(),
                })
            continue
        periodo = str(cell(row, mapping, "periodo") or "").strip().upper()
        processo = str(cell(row, mapping, "processo") or "").strip()
        if surgery_date < today:
            anomalies.append({"type": "past_surgery_ignored", "row": excel_row_idx, "date": surgery_date.isoformat(), "processo": processo})
            continue
        if surgery_date > end:
            continue
        if periodo not in {"M", "S"}:
            anomalies.append({"type": "invalid_periodo", "row": excel_row_idx, "date": surgery_date.isoformat(), "periodo": periodo})
            periodo = periodo or "M"
        if not processo:
            anomalies.append({"type": "missing_processo", "row": excel_row_idx, "date": surgery_date.isoformat(), "periodo": periodo})
        if periodo == "M" and surgery_date.weekday() not in {0, 2}:
            anomalies.append({"type": "manha_on_non_bo_weekday", "row": excel_row_idx, "date": surgery_date.isoformat(), "weekday": WEEKDAYS_PT[surgery_date.weekday()]})

        ausentes = []
        if parse_bool(cell(row, mapping, "if_ausente")):
            ausentes.append("IF")
        if parse_bool(cell(row, mapping, "rc_ausente")):
            ausentes.append("RC")
        if parse_bool(cell(row, mapping, "rr_ausente")):
            ausentes.append("RR")

        patient = {
            "processo": processo,
            "nome": str(cell(row, mapping, "nome") or "").strip(),
            "idade_raw": parse_float(cell(row, mapping, "idade")),
            "idade_fmt": age_fmt(cell(row, mapping, "idade")),
            "procedimento": str(cell(row, mapping, "procedimento") or "").strip(),
            "fdr": parse_bool(cell(row, mapping, "fdr")),
            "tempo_espera": parse_float(cell(row, mapping, "tempo_espera")),
            "obs": str(cell(row, mapping, "obs") or "").strip(),
            "source_row": excel_row_idx,
        }
        raw_rows.append(
            {
                "date": surgery_date.isoformat(),
                "weekday": WEEKDAYS_PT[surgery_date.weekday()],
                "periodo": periodo,
                "processo": processo,
                "fdr": patient["fdr"],
                "ausentes_previstos": ausentes,
                "source_row": excel_row_idx,
            }
        )
        if processo:
            process_keys.append((surgery_date.isoformat(), periodo, processo))
        key = (surgery_date.isoformat(), periodo)
        session = sessions.setdefault(
            key,
            {
                "date": surgery_date.isoformat(),
                "weekday": WEEKDAYS_PT[surgery_date.weekday()],
                "periodo": periodo,
                "cirurgioes": [],
                "ausentes_previstos": [],
                "ausencias_provaveis": [],
                "doentes": [],
            },
        )
        for initials in ausentes:
            if initials not in session["ausentes_previstos"]:
                session["ausentes_previstos"].append(initials)
                session["ausencias_provaveis"].append({"cir": initials, "fonte": "flag_cirurgias"})
        session["doentes"].append(patient)

    duplicated = [list(key) for key, count in Counter(process_keys).items() if count > 1]
    for key in duplicated:
        anomalies.append({"type": "duplicate_processo_same_session", "date": key[0], "periodo": key[1], "processo": key[2]})

    ordered_sessions = []
    for session in sessions.values():
        session["ausentes_previstos"] = sorted(session["ausentes_previstos"])
        session["ausencias_provaveis"] = sorted(session["ausencias_provaveis"], key=lambda item: item["cir"])
        session["doentes"] = sorted(session["doentes"], key=lambda item: (item["idade_raw"] is None, item["idade_raw"] or 0, item["nome"]))
        ordered_sessions.append(session)
    ordered_sessions.sort(key=lambda item: (item["date"], item["periodo"]))

    return {
        "ok": True,
        "headers_found": sorted(mapping),
        "missing_headers": missing_headers,
        "window": {"start": today.isoformat(), "end": end.isoformat(), "days": WINDOW_DAYS},
        "rows_in_window": len(raw_rows),
        "sessions_count": len(ordered_sessions),
        "fdr_count": sum(1 for row in raw_rows if row["fdr"]),
        "duplicates": duplicated,
        "anomalies": anomalies,
        "sessions": ordered_sessions,
        "raw_rows": raw_rows,
    }


def parse_sigic(ws: Any, today: date, warnings: list[str], errors: list[str]) -> dict[str, Any]:
    aliases = {
        "mes": ["MES", "Mes"],
        "dia": ["DIA de Sigic", "Dia de SIGIC", "Data", "Dia"],
        "cirurgiao1": ["Cirurgiao1", "Cirurgião1", "Cirurgiao 1", "Cirurgião 1"],
        "cirurgiao2": ["Cirurgiao2", "Cirurgião2", "Cirurgiao 2", "Cirurgião 2"],
    }
    rows = row_values(ws)
    header_idx, mapping = find_header(rows, aliases, minimum=3)
    if header_idx is None or len(mapping) < 3:
        errors.append("SIGIC: cabecalho obrigatorio nao encontrado")
        return {"ok": False, "headers_found": sorted(mapping)}

    end = today + timedelta(days=WINDOW_DAYS - 1)
    lists = []
    anomalies = []
    for excel_row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        dia_value = cell(row, mapping, "dia")
        sigic_date = parse_date(dia_value, default_year=today.year)
        if sigic_date is None and isinstance(dia_value, (int, float)):
            month = month_from_value(cell(row, mapping, "mes"))
            if month:
                try:
                    sigic_date = date(today.year, month, int(dia_value))
                except ValueError:
                    sigic_date = None
        if sigic_date is None:
            continue
        if not (today <= sigic_date <= end):
            continue
        item = {
            "date": sigic_date.isoformat(),
            "cirurgiao1": str(cell(row, mapping, "cirurgiao1") or "").strip(),
            "cirurgiao2": str(cell(row, mapping, "cirurgiao2") or "").strip(),
            "cirurgias": [],
            "source_row": excel_row_idx,
        }
        if not item["cirurgiao1"] or not item["cirurgiao2"]:
            anomalies.append({"type": "missing_sigic_surgeon", "row": excel_row_idx, "date": item["date"]})
        lists.append(item)
    lists.sort(key=lambda item: item["date"])
    return {"ok": True, "headers_found": sorted(mapping), "listas": lists, "anomalies": anomalies}


def month_from_value(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and 1 <= int(value) <= 12:
        return int(value)
    text = norm(value)
    if text in MONTHS_PT:
        return MONTHS_PT[text]
    return None


def parse_absences_best_effort(ws: Any, today: date) -> dict[str, Any]:
    rows = row_values(ws)
    end = today + timedelta(days=WINDOW_DAYS - 1)
    absences: list[dict[str, Any]] = []
    day_offs: list[dict[str, Any]] = []
    ambiguous: list[dict[str, Any]] = []
    undated_codes: list[dict[str, Any]] = []

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row, start=1):
            text = str(value or "").strip()
            if not text:
                continue
            compact = compact_norm(text)
            absence_match = ABSENCE_RE.match(compact)
            day_off_match = DAY_OFF_RE.match(compact)
            ambiguous_single = compact.upper() in {"F", "C", "M"}
            if not (absence_match or day_off_match or ambiguous_single):
                continue
            parsed_date = infer_grid_date(rows, r_idx - 1, c_idx - 1, today.year)
            record = {"codigo": text, "source_cell": f"{ws.cell(r_idx, c_idx).coordinate}"}
            if parsed_date is None:
                undated_codes.append(record)
                continue
            if not (today <= parsed_date <= end):
                continue
            record["data"] = parsed_date.isoformat()
            if absence_match:
                code = absence_match.group("motive").upper()
                initials = absence_match.group("who").upper()
                absences.append({"cir": initials, "motivo": motive_label(code), **record})
            elif day_off_match:
                initials = day_off_match.group("who").upper()
                day_offs.append({"cir": initials, **record})
            else:
                ambiguous.append({"nota": "Codigo sem iniciais; nao atribuido.", **record})

    return {
        "ok": True,
        "ausencias": sorted(absences, key=lambda item: (item["data"], item["cir"])),
        "folgas": sorted(day_offs, key=lambda item: (item["data"], item["cir"])),
        "codigos_ambiguos": sorted(ambiguous, key=lambda item: item["data"]),
        "undated_codes": undated_codes,
        "parser_note": "Leitura best-effort: confirmar layout real da grelha se undated_codes > 0.",
    }


def infer_grid_date(rows: list[list[Any]], r: int, c: int, year: int) -> date | None:
    explicit = parse_date(rows[r][c], default_year=year)
    if explicit:
        return explicit

    day: int | None = None
    for rr, cc in ((r, c - 1), (r, c - 2), (r - 1, c), (r - 2, c), (r - 1, c - 1)):
        if rr < 0 or cc < 0:
            continue
        try:
            candidate = rows[rr][cc]
        except IndexError:
            continue
        parsed = parse_date(candidate, default_year=year)
        if parsed:
            return parsed
        if isinstance(candidate, (int, float)) and 1 <= int(candidate) <= 31:
            day = int(candidate)
            break
        if isinstance(candidate, str) and re.fullmatch(r"\d{1,2}", candidate.strip()):
            day = int(candidate.strip())
            break
    if day is None:
        return None

    month = None
    for rr in range(max(0, r - 8), min(len(rows), r + 3)):
        for cc in range(max(0, c - 8), min(len(rows[rr]), c + 3)):
            month = month_from_value(rows[rr][cc])
            if month:
                break
        if month:
            break
    if not month:
        return None
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_prevencao(ws: Any, today: date, absence_data: dict[str, Any], warnings: list[str], errors: list[str]) -> dict[str, Any]:
    aliases = {
        "semana": ["Semana"],
        "inicio": ["Data Inicio", "Data Início", "Inicio", "Início"],
        "fim": ["Data Fim", "Data fim", "Fim"],
        "cirurgiao": ["Cirurgiao de prevencao", "Cirurgião de prevenção", "Cirurgiao a fazer a Prevencao", "Cirurgião a fazer a Prevenção"],
        "madeira": ["Cirurgiao na Madeira", "Cirurgião na Madeira", "Cirurgiao com alguns dias da semana na Madeira (Nao pode estar de prevencao)"],
        "ausente": ["Cirurgiao de ferias/ausente", "Cirurgião de férias/ausente", "Cirurgiao de ferias ou ausente por outro motivo"],
        "obs": ["Observacoes", "Observações"],
    }
    rows = row_values(ws)
    header_idx, mapping = find_header(rows, aliases, minimum=4)
    if header_idx is None or len(mapping) < 4:
        errors.append("Prevencao: cabecalho obrigatorio nao encontrado")
        return {"ok": False, "headers_found": sorted(mapping)}

    month_start = date(today.year, today.month, 1)
    weeks = []
    for excel_row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        cirurgiao = str(cell(row, mapping, "cirurgiao") or "").strip()
        if not cirurgiao:
            continue
        start = parse_date(cell(row, mapping, "inicio"), default_year=today.year)
        finish = parse_date(cell(row, mapping, "fim"), default_year=today.year)
        if not start or not finish:
            continue
        if finish < month_start:
            continue
        notes = build_prevention_absence_notes(row, mapping, start, finish, absence_data)
        weeks.append(
            {
                "semana": cell(row, mapping, "semana"),
                "inicio": start.isoformat(),
                "fim": finish.isoformat(),
                "cirurgiao": cirurgiao,
                "ausencia_notas": notes,
                "obs": str(cell(row, mapping, "obs") or "").strip(),
                "source_row": excel_row_idx,
            }
        )
    weeks.sort(key=lambda item: item["inicio"])
    current = next((week for week in weeks if week["inicio"] <= today.isoformat() <= week["fim"]), None)
    return {
        "ok": True,
        "headers_found": sorted(mapping),
        "semana_actual": {
            "cirurgiao": (current or {}).get("cirurgiao", ""),
            "on_call_rafael": "rafael" in norm((current or {}).get("cirurgiao", "")),
        },
        "semanas": weeks,
    }


def build_prevention_absence_notes(
    row: list[Any],
    mapping: dict[str, int],
    start: date,
    finish: date,
    absence_data: dict[str, Any],
) -> list[str]:
    notes: list[str] = []
    madeira = str(cell(row, mapping, "madeira") or "").strip()
    generic_absent = str(cell(row, mapping, "ausente") or "").strip()
    if madeira:
        notes.append(f"{madeira}: Madeira esta semana")
    initials = initials_from_name(generic_absent)
    if generic_absent and initials:
        matching = [
            item
            for item in absence_data.get("ausencias", [])
            if item.get("cir") == initials and start.isoformat() <= item.get("data", "") <= finish.isoformat()
        ]
        if matching:
            dates = ", ".join(item["data"] for item in matching)
            motives = sorted({item["motivo"] for item in matching})
            notes.append(f"{generic_absent}: {'/'.join(motives)} ({dates})")
        else:
            notes.append(f"{generic_absent}: ausencia parte desta semana; motivo nao especificado na aba Ausencias")
    elif generic_absent:
        notes.append(f"{generic_absent}: ausencia parte desta semana; pessoa nao reconhecida mecanicamente")
    return notes


def parse_todoist_json(path: Path | None, today: date, warnings: list[str]) -> dict[str, Any]:
    if path is None:
        return {
            "ok": False,
            "required_source": "Todoist directo",
            "note": "Fornecer --todoist-json com tarefas directas Todoist; calendario Todoist nao contem prioridade.",
            "itens": [],
        }
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        warnings.append(f"Todoist: JSON invalido: {type(exc).__name__}: {exc}")
        return {"ok": False, "path": str(path), "itens": []}
    items = data if isinstance(data, list) else data.get("tasks", data.get("items", [])) if isinstance(data, dict) else []
    output = []
    end = today + timedelta(days=2)
    for item in items if isinstance(items, list) else []:
        if not isinstance(item, dict):
            continue
        priority_raw = item.get("priority", item.get("prioridade"))
        priority = classify_priority(priority_raw)
        if priority not in {"alta", "media"}:
            continue
        due_value = item.get("date") or item.get("due_date")
        due = item.get("due")
        if due_value is None and isinstance(due, dict):
            due_value = due.get("date")
        due_date = parse_date(due_value, default_year=today.year)
        if due_date and not (today <= due_date <= end):
            continue
        output.append(
            {
                "texto": str(item.get("content") or item.get("texto") or item.get("title") or "").strip(),
                "prioridade": priority,
                "date": due_date.isoformat() if due_date else today.isoformat(),
                "prazo": "hoje" if priority == "alta" else "ate_2_dias",
            }
        )
    return {"ok": True, "path": str(path), "itens": output}


def classify_priority(value: Any) -> str:
    text = norm(value)
    if text in {"alta", "high", "p1", "4"}:
        return "alta"
    if text in {"media", "média", "medium", "p2", "3"}:
        return "media"
    if isinstance(value, int):
        if value >= 4:
            return "alta"
        if value == 3:
            return "media"
    return "baixa"


def attach_sigic_to_sessions(cirurgias: dict[str, Any], sigic: dict[str, Any]) -> None:
    sigic_by_date = {item["date"]: item for item in sigic.get("listas", [])}
    for session in cirurgias.get("sessions", []):
        if session.get("periodo") != "S":
            continue
        listing = sigic_by_date.get(session.get("date"))
        if listing:
            session["cirurgioes"] = [listing.get("cirurgiao1", ""), listing.get("cirurgiao2", "")]
            listing["cirurgias"] = session.get("doentes", [])


def build_day_type(today: date, sessions: list[dict[str, Any]], sigic: dict[str, Any]) -> str:
    has_m = any(session["date"] == today.isoformat() and session["periodo"] == "M" for session in sessions)
    has_s = any(session["date"] == today.isoformat() and session["periodo"] == "S" for session in sessions)
    has_sigic_sheet = any(item["date"] == today.isoformat() for item in sigic.get("listas", []))
    if has_m and (has_s or has_sigic_sheet):
        return "BO manha + SIGIC tarde"
    if has_m:
        return "BO manha"
    if has_s or has_sigic_sheet:
        return "SIGIC tarde"
    return "Sem BO"


def build_resumo(today: date, cirurgias: dict[str, Any], absence_data: dict[str, Any], day_type: str) -> dict[str, Any]:
    sessions = cirurgias.get("sessions", [])
    alertas: list[str] = []
    if cirurgias.get("fdr_count"):
        alertas.append(f"{cirurgias['fdr_count']} doente(s) FDR na janela HFF")
    if cirurgias.get("duplicates"):
        alertas.append(f"{len(cirurgias['duplicates'])} processo(s) duplicado(s) na mesma sessao")
    if absence_data.get("codigos_ambiguos"):
        alertas.append(f"{len(absence_data['codigos_ambiguos'])} codigo(s) de ausencia ambiguo(s)")

    future_bo = []
    for session in sessions:
        if session["periodo"] != "M":
            continue
        if session["date"] >= today.isoformat():
            future_bo.append((session["date"], len(session.get("doentes", []))))
    proximo_bo = None
    if future_bo:
        next_date, n_surg = sorted(future_bo)[0]
        dt = date.fromisoformat(next_date)
        proximo_bo = {"date": next_date, "weekday": WEEKDAYS_PT[dt.weekday()], "n_cirurgias": n_surg}

    return {
        "tipo_dia": day_type,
        "alertas": alertas,
        "ausencias_relevantes": summarize_absences_for_resumo(absence_data),
        "proximo_bo": proximo_bo,
    }


def summarize_absences_for_resumo(absence_data: dict[str, Any]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], list[str]] = defaultdict(list)
    for item in absence_data.get("ausencias", []):
        grouped[(item["cir"], item["motivo"])].append(item["data"])
    for item in absence_data.get("folgas", []):
        grouped[(item["cir"], "Folga")].append(item["data"])
    output = []
    for (cir, motive), dates in sorted(grouped.items()):
        output.append({"cir": cir, "motivo": motive, "datas": compact_date_list(sorted(dates))})
    return output


def compact_date_list(values: list[str]) -> str:
    if not values:
        return ""
    if len(values) == 1:
        return values[0]
    return f"{values[0]} a {values[-1]}"


def build_equipa(today: date, absence_data: dict[str, Any], cirurgias: dict[str, Any], sigic: dict[str, Any]) -> dict[str, Any]:
    sessions = cirurgias.get("sessions", [])
    confirmed = summarize_absences_for_resumo({"ausencias": absence_data.get("ausencias", []), "folgas": []})
    probable = []
    for session in sessions:
        for initials in session.get("ausentes_previstos", []):
            probable.append({"cir": initials, "dia": session["date"], "fonte": "flag_cirurgias"})
    sigic_by_date = {item["date"]: {"cirurgiao1": item.get("cirurgiao1", ""), "cirurgiao2": item.get("cirurgiao2", "")} for item in sigic.get("listas", [])}
    morning_by_date = defaultdict(int)
    for session in sessions:
        if session.get("periodo") == "M":
            morning_by_date[session["date"]] += len(session.get("doentes", []))
    abs_by_date = defaultdict(list)
    for item in absence_data.get("ausencias", []):
        abs_by_date[item["data"]].append({"cir": item["cir"], "motivo": item["motivo"]})
    calendar = []
    for offset in range(WINDOW_DAYS):
        day = today + timedelta(days=offset)
        key = day.isoformat()
        calendar.append(
            {
                "date": key,
                "weekday": WEEKDAYS_PT[day.weekday()],
                "ausencias": sorted(abs_by_date.get(key, []), key=lambda item: item["cir"]),
                "n_cirurgias_manha": morning_by_date.get(key, 0),
                "sigic_tarde": sigic_by_date.get(key),
            }
        )
    return {
        "ausencias_confirmadas": confirmed,
        "ausencias_provaveis": sorted(probable, key=lambda item: (item["dia"], item["cir"])),
        "folgas": [{"cir": item["cir"], "data": item["data"]} for item in absence_data.get("folgas", [])],
        "codigos_ambiguos": absence_data.get("codigos_ambiguos", []),
        "calendario_4_semanas": calendar,
    }


def build_canonical_draft(
    today: date,
    generated_at: str,
    cirurgias: dict[str, Any],
    sigic: dict[str, Any],
    prevencao: dict[str, Any],
    absence_data: dict[str, Any],
    todoist: dict[str, Any],
) -> dict[str, Any]:
    attach_sigic_to_sessions(cirurgias, sigic)
    day_type = build_day_type(today, cirurgias.get("sessions", []), sigic)
    return {
        "meta": {
            "date": today.isoformat(),
            "weekday": WEEKDAYS_PT[today.weekday()],
            "day_type": day_type,
            "generated_at": generated_at,
        },
        "resumo": build_resumo(today, cirurgias, absence_data, day_type),
        "tarefas_hff": {
            "janela": "hoje + 2 dias",
            "itens": todoist.get("itens", []),
        },
        "cirurgias": {
            "janela_dias": WINDOW_DAYS,
            "sessoes": cirurgias.get("sessions", []),
        },
        "prevencao": {
            "semana_actual": prevencao.get("semana_actual", {"cirurgiao": "", "on_call_rafael": False}),
            "semanas": prevencao.get("semanas", []),
        },
        "sigic": {
            "listas": sigic.get("listas", []),
        },
        "equipa": build_equipa(today, absence_data, cirurgias, sigic),
    }


def inspect_workbook(path: Path, today: date, warnings: list[str], errors: list[str]) -> dict[str, Any]:
    try:
        wb = load_workbook(path)
    except Exception as exc:
        errors.append(f"Espelho HFF: impossivel abrir xlsx: {type(exc).__name__}: {exc}")
        return {"ok": False, "file": file_info(path)}

    sheets = {
        "cirurgias": sheet_by_name(wb, ["Cirurgias"]),
        "sigic": sheet_by_name(wb, ["SIGIC"]),
        "prevencao": sheet_by_name(wb, ["Prevencao", "Prevenção"]),
        "ausencias": sheet_by_name(wb, ["Ausencias", "Ausências"]),
    }
    for key, ws in sheets.items():
        if ws is None:
            errors.append(f"Espelho HFF: aba obrigatoria em falta: {key}")

    absence_data = parse_absences_best_effort(sheets["ausencias"], today) if sheets["ausencias"] is not None else {"ok": False, "ausencias": [], "folgas": [], "codigos_ambiguos": []}
    cirurgias = parse_cirurgias(sheets["cirurgias"], today, warnings, errors) if sheets["cirurgias"] is not None else {"ok": False, "sessions": []}
    sigic = parse_sigic(sheets["sigic"], today, warnings, errors) if sheets["sigic"] is not None else {"ok": False, "listas": []}
    prevencao = parse_prevencao(sheets["prevencao"], today, absence_data, warnings, errors) if sheets["prevencao"] is not None else {"ok": False, "semanas": []}
    if absence_data.get("undated_codes"):
        warnings.append(f"Ausencias: {len(absence_data['undated_codes'])} codigo(s) encontrados sem data inferida")

    return {
        "ok": not any(ws is None for ws in sheets.values()),
        "file": file_info(path),
        "sheetnames": wb.sheetnames,
        "sheets_present": {key: ws is not None for key, ws in sheets.items()},
        "cirurgias": cirurgias,
        "sigic": sigic,
        "prevencao": prevencao,
        "ausencias": absence_data,
    }


def inspect_sheets_values(path: Path, today: date, warnings: list[str], errors: list[str]) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        errors.append(f"Sheets JSON: impossivel ler: {type(exc).__name__}: {exc}")
        return {"ok": False, "file": file_info(path)}

    sheets_payload = payload.get("sheets") if isinstance(payload, dict) else None
    if not isinstance(sheets_payload, dict):
        errors.append("Sheets JSON: campo 'sheets' deve ser objecto")
        return {"ok": False, "file": file_info(path)}

    wb = ValueWorkbook({str(name): rows for name, rows in sheets_payload.items() if isinstance(rows, list)})
    sheets = {
        "cirurgias": sheet_by_name(wb, ["Cirurgias"]),
        "sigic": sheet_by_name(wb, ["SIGIC", "Sigic"]),
        "prevencao": sheet_by_name(wb, ["Prevencao", "Prevenção"]),
        "ausencias": sheet_by_name(wb, ["Ausencias", "Ausências"]),
    }
    for key, ws in sheets.items():
        if ws is None:
            errors.append(f"Sheets JSON: aba obrigatoria em falta: {key}")

    absence_data = parse_absences_best_effort(sheets["ausencias"], today) if sheets["ausencias"] is not None else {"ok": False, "ausencias": [], "folgas": [], "codigos_ambiguos": []}
    cirurgias = parse_cirurgias(sheets["cirurgias"], today, warnings, errors) if sheets["cirurgias"] is not None else {"ok": False, "sessions": []}
    sigic = parse_sigic(sheets["sigic"], today, warnings, errors) if sheets["sigic"] is not None else {"ok": False, "listas": []}
    prevencao = parse_prevencao(sheets["prevencao"], today, absence_data, warnings, errors) if sheets["prevencao"] is not None else {"ok": False, "semanas": []}
    if absence_data.get("undated_codes"):
        warnings.append(f"Ausencias: {len(absence_data['undated_codes'])} codigo(s) encontrados sem data inferida")

    return {
        "ok": not any(ws is None for ws in sheets.values()),
        "file": file_info(path),
        "sheetnames": wb.sheetnames,
        "sheets_present": {key: ws is not None for key, ws in sheets.items()},
        "cirurgias": cirurgias,
        "sigic": sigic,
        "prevencao": prevencao,
        "ausencias": absence_data,
    }


def build_manifest(args: argparse.Namespace) -> dict[str, Any]:
    generated = now_lisbon()
    today = date.fromisoformat(args.today) if args.today else generated.date()
    warnings: list[str] = []
    errors: list[str] = []
    templates = inspect_templates_and_repo()

    for key in ("instrucoes_hff", "rotina_publicacao", "system_prompt", "hff_template", "hff_render", "bo_template", "bo_render"):
        if not templates[key]["exists"]:
            errors.append(f"Ficheiro obrigatorio em falta: {key}")
    if not templates["repo"]["exists"]:
        warnings.append("Checkout local personal-system nao encontrado; postflight nao conseguira publicar localmente")

    todoist_path = Path(args.todoist_json) if args.todoist_json else None
    todoist = parse_todoist_json(todoist_path, today, warnings)
    workbook = None
    canonical_draft = None
    if args.sheets_json:
        workbook = inspect_sheets_values(Path(args.sheets_json), today, warnings, errors)
        canonical_draft = build_canonical_draft(
            today=today,
            generated_at=generated.isoformat(timespec="seconds"),
            cirurgias=workbook.get("cirurgias", {}),
            sigic=workbook.get("sigic", {}),
            prevencao=workbook.get("prevencao", {}),
            absence_data=workbook.get("ausencias", {}),
            todoist=todoist,
        )
    elif args.xlsx:
        workbook = inspect_workbook(Path(args.xlsx), today, warnings, errors)
        canonical_draft = build_canonical_draft(
            today=today,
            generated_at=generated.isoformat(timespec="seconds"),
            cirurgias=workbook.get("cirurgias", {}),
            sigic=workbook.get("sigic", {}),
            prevencao=workbook.get("prevencao", {}),
            absence_data=workbook.get("ausencias", {}),
            todoist=todoist,
        )
    elif not args.allow_missing_xlsx:
        errors.append("Espelho HFF nao fornecido. Exportar xlsx e chamar --xlsx <ficheiro>.")
    else:
        warnings.append("Espelho HFF nao fornecido; manifesto limitado a ficheiros locais")

    anomalies = []
    if workbook:
        anomalies.extend(workbook.get("cirurgias", {}).get("anomalies", []))
        anomalies.extend(workbook.get("sigic", {}).get("anomalies", []))

    status = "fail" if errors else ("warn" if warnings or anomalies else "ok")
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": generated.isoformat(timespec="seconds"),
        "today": today.isoformat(),
        "status": status,
        "routine": "painel_hff",
        "inputs": {
            "espelho_hff": {
                "spreadsheet_id": ESPELHO_HFF_ID,
                "xlsx": file_info(Path(args.xlsx)) if args.xlsx else {"exists": False, "path": None},
                "sheets_json": file_info(Path(args.sheets_json)) if args.sheets_json else {"exists": False, "path": None},
            },
            "todoist": {
                **file_info(todoist_path),
                "parsed": todoist,
            },
            "local_files": templates,
        },
        "workbook": workbook,
        "manifesto_cobertura": {
            "window_days": WINDOW_DAYS,
            "cirurgias_sessions": len((workbook or {}).get("cirurgias", {}).get("sessions", [])),
            "cirurgias_rows": (workbook or {}).get("cirurgias", {}).get("rows_in_window"),
            "sigic_lists": len((workbook or {}).get("sigic", {}).get("listas", [])),
            "prevencao_weeks": len((workbook or {}).get("prevencao", {}).get("semanas", [])),
            "ausencias_confirmadas": len((workbook or {}).get("ausencias", {}).get("ausencias", [])),
            "folgas": len((workbook or {}).get("ausencias", {}).get("folgas", [])),
            "codigos_ambiguos": len((workbook or {}).get("ausencias", {}).get("codigos_ambiguos", [])),
            "anomalies": anomalies,
            "critical": bool(errors),
        },
        "canonical_draft": canonical_draft,
        "llm_handoff": {
            "role": "conferencia_inteligente",
            "instructions": [
                "Nao refazer mecanicamente a planilha; usar canonical_draft como base.",
                "Rever alertas, lacunas, FDR e prioridades clinicas/operacionais.",
                "Se manifesto_cobertura.critical=true, reportar bloqueio antes de redigir JSON final.",
                "Emitir apenas JSON canonico final para o postflight.",
            ],
            "nunca_reescrever": [
                "cirurgias (sessoes, doentes, ausentes_previstos) -- ja vem apurado da planilha",
                "sigic.listas -- idem",
                "equipa (ausencias_confirmadas, ausencias_provaveis, folgas, codigos_ambiguos, "
                "calendario_4_semanas) -- ja vem apurado da aba Ausencias, nunca inferir/assumir motivo",
                "prevencao.semanas[].ausencia_notas -- ja vem cruzado com a aba Ausencias (Seccao 7-A); "
                "nunca assumir 'ferias' quando a coluna ambigua da Prevencao tiver um valor",
                "meta.date/weekday/day_type",
            ],
        },
        "warnings": warnings,
        "errors": errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Preflight mecanico da rotina Painel HFF.")
    parser.add_argument("--xlsx", help="Export xlsx do Espelho HFF.")
    parser.add_argument("--sheets-json", help="JSON de valores lidos pelo conector Google Sheets.")
    parser.add_argument("--todoist-json", help="JSON opcional de tarefas Todoist directas.")
    parser.add_argument("--today", help="Data alvo YYYY-MM-DD; por omissao, data real Europe/Lisbon.")
    parser.add_argument("--out", default=str(ROUTINE_DIR / "hff_preflight.json"), help="Manifesto JSON de saida.")
    parser.add_argument("--allow-missing-xlsx", action="store_true", help="Permite auditoria parcial sem xlsx.")
    args = parser.parse_args()

    manifest = build_manifest(args)
    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": manifest["status"], "manifest": str(out_path)}, ensure_ascii=False))
    return 0 if manifest["status"] != "fail" else 1


if __name__ == "__main__":
    raise SystemExit(main())
